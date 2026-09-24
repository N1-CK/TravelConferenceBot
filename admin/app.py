import asyncio
import csv
import io
import logging
import mimetypes
import os
import secrets
import sys
import threading
import concurrent.futures
import uuid
from datetime import datetime, timedelta
from functools import wraps
from urllib.parse import unquote

import openpyxl
import requests
from aiogram.client.session import aiohttp
from dotenv import load_dotenv
from flask import Flask, flash, jsonify, redirect, render_template, request, send_file, session, url_for
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from werkzeug.utils import secure_filename

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from database import db
from locales import get_text
from utility.notifications import (
    notify_banner_status_change,
    notify_business_card_status_change,
    notify_ticket_status_change,
    notify_travel_flight_status_change,
    notify_travel_visa_status_change,
)

logger = logging.getLogger(__name__)

load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv('FLASK_SECRET_KEY', secrets.token_hex(32))
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max file size
app.config['UPLOAD_FOLDER'] = '/tmp'

FORBIDDEN_EXTENSIONS = {
    'exe', 'bat', 'cmd', 'sh', 'bash', 'bin', 'msi', 'vbs', 'ps1', 'jar', 'apk', 'com', 'scr'
}
STANDS_UPLOAD_FOLDER = os.path.join(app.root_path, 'static', 'uploads', 'stands')
os.makedirs(STANDS_UPLOAD_FOLDER, exist_ok=True)
ALLOWED_IMAGE_EXTENSIONS = {'png', 'jpg', 'jpeg', 'webp'}

def allowed_image(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_IMAGE_EXTENSIONS

def allowed_file(filename):
    if '.' not in filename:
        return True
    ext = filename.rsplit('.', 1)[1].lower()
    return ext not in FORBIDDEN_EXTENSIONS

# ============================================
# EVENT LOOP ДЛЯ ФОНОВЫХ ЗАДАЧ
# ============================================

_bg_loop = asyncio.new_event_loop()
_db_initialized = False

def _run_bg_loop(loop):
    asyncio.set_event_loop(loop)
    loop.run_forever()

_bg_thread = threading.Thread(target=_run_bg_loop, args=(_bg_loop,), daemon=True)
_bg_thread.start()

@app.context_processor
def inject_request():
    return {'request': request}

def run_async(coro, timeout=30):
    future = asyncio.run_coroutine_threadsafe(coro, _bg_loop)
    try:
        return future.result(timeout=timeout)
    except concurrent.futures.TimeoutError:
        logger.error("Async operation timeout")
        return None
    except Exception as e:
        logger.error(f"Async execution error: {e}")
        return None

def init_db():
    global _db_initialized
    if not _db_initialized:
        try:
            run_async(db.create_pool())
            run_async(db.create_managers_tables())
            _db_initialized = True
            print("✅ Database connection initialized")
            print("✅ Managers tables created")
        except Exception as e:
            print(f"❌ Database connection failed: {e}")
            raise

init_db()

# ============================================
# ДЕКОРАТОРЫ АВТОРИЗАЦИИ
# ============================================

def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('logged_in'):
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function

def group_required(allowed_groups=None):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if not session.get('logged_in'):
                return redirect(url_for('login'))

            if session.get('role') == 'admin':
                return f(*args, **kwargs)

            user_groups = session.get('groups', [])
            if allowed_groups is None:
                return f(*args, **kwargs)

            if any(group in user_groups for group in allowed_groups):
                return f(*args, **kwargs)

            flash('У вас нет доступа к этой странице', 'danger')
            return redirect(url_for('dashboard'))
        return decorated_function
    return decorator

def role_required(allowed_roles=None, permissions=None):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if not session.get('logged_in'):
                return redirect(url_for('login'))

            role = session.get('role', 'user')
            if allowed_roles and role not in allowed_roles:
                flash('У вас нет доступа к этой странице', 'danger')
                return redirect(url_for('dashboard'))

            if permissions:
                for perm in permissions:
                    if not session.get('permissions', {}).get(perm, False):
                        flash(f'У вас нет права "{perm}"', 'danger')
                        return redirect(url_for('dashboard'))
            return f(*args, **kwargs)
        return decorated_function
    return decorator

# ============================================
# ИНТЕГРАЦИЯ С TELEGRAM БОТОМ
# ============================================

class TelegramBot:
    def __init__(self, bot_token):
        self.bot_token = bot_token
        self.api_url = f"https://api.telegram.org/bot{bot_token}"

    async def send_message(self, chat_id, text, file=None):
        import aiohttp
        timeout = aiohttp.ClientTimeout(total=300, connect=20, sock_read=300)
        async with aiohttp.ClientSession(timeout=timeout) as session:
            if file:
                ext = file.rsplit('.', 1)[-1].lower() if '.' in file else ''
                if ext in {'ogg', 'oga', 'opus'}:
                    url = f"{self.api_url}/sendVoice"
                    field_name = 'voice'
                elif ext in {'mp3', 'm4a', 'wav', 'flac'}:
                    url = f"{self.api_url}/sendAudio"
                    field_name = 'audio'
                elif ext in {'jpg', 'jpeg', 'png', 'webp'} and os.path.getsize(file) <= 10 * 1024 * 1024:
                    url = f"{self.api_url}/sendPhoto"
                    field_name = 'photo'
                else:
                    url = f"{self.api_url}/sendDocument"
                    field_name = 'document'

                form_data = aiohttp.FormData()
                form_data.add_field('chat_id', str(chat_id))
                with open(file, 'rb') as f:
                    form_data.add_field(field_name, f, filename=os.path.basename(file))
                    if text:
                        form_data.add_field('caption', text)
                        form_data.add_field('parse_mode', 'HTML')
                    async with session.post(url, data=form_data) as resp:
                        result = await resp.json()
                        if result.get('ok'):
                            res = result.get('result', {})
                            doc = res.get('document') or res.get('voice') or res.get('audio') or (
                                res.get('photo', [{}])[-1] if 'photo' in res else {})
                            return doc.get('file_id') or True
                        else:
                            logger.error(f"Telegram API send error: {result}")
                            return False
            else:
                url = f"{self.api_url}/sendMessage"
                payload = {'chat_id': chat_id, 'text': text, 'parse_mode': 'HTML'}
                async with session.post(url, json=payload) as resp:
                    result = await resp.json()
                    return result.get('ok', False)

    async def send_documents(self, chat_id, text, files=None):
        import aiohttp
        import json

        if not files:
            await self.send_message(chat_id, text)
            return []

        if len(files) == 1:
            file_id = await self.send_message(chat_id, text, files[0])
            return [file_id] if isinstance(file_id, str) else []

        timeout = aiohttp.ClientTimeout(total=300, connect=30, sock_read=300)
        form_data = aiohttp.FormData()
        form_data.add_field('chat_id', str(chat_id))

        media_group = []
        opened_files = []
        try:
            for idx, file_path in enumerate(files):
                attach_name = f"file_{idx}"
                f = open(file_path, 'rb')
                opened_files.append(f)
                form_data.add_field(attach_name, f, filename=os.path.basename(file_path))

                ext = file_path.rsplit('.', 1)[-1].lower() if '.' in file_path else ''
                if ext in {'jpg', 'jpeg', 'png'} and os.path.getsize(file_path) <= 10 * 1024 * 1024:
                    media_type = 'photo'
                elif ext in {'mp3', 'm4a', 'wav', 'flac'}:
                    media_type = 'audio'
                else:
                    media_type = 'document'

                item = {'type': media_type, 'media': f'attach://{attach_name}'}
                if idx == 0 and text:
                    item['caption'] = text
                    item['parse_mode'] = 'HTML'
                media_group.append(item)

            form_data.add_field('media', json.dumps(media_group))

            async with aiohttp.ClientSession(timeout=timeout) as session:
                url = f"{self.api_url}/sendMediaGroup"
                async with session.post(url, data=form_data) as resp:
                    res_json = await resp.json()
                    file_ids = []
                    if res_json.get('ok'):
                        for msg in res_json.get('result', []):
                            doc = msg.get('document') or msg.get('voice') or msg.get('audio') or (
                                msg.get('photo', [{}])[-1] if 'photo' in msg else {})
                            if 'file_id' in doc:
                                file_ids.append(doc['file_id'])
                        return file_ids
                    else:
                        logger.error(f"Telegram error sendMediaGroup: {res_json}")
                        return []
        finally:
            for f in opened_files:
                f.close()

    async def broadcast_to_users(self, users, message, files=None):
        results = {
            'success': 0, 'failed': 0, 'file_ids': [], 'success_users': [], 'failed_users': []
        }
        file_count = len(files) if files else 0
        file_ids_map = {}

        for user in users:
            user_id = user.get('user_id')
            username = user.get('username') or 'unknown'
            full_name = user.get('full_name') or ''
            user_info = {'user_id': user_id, 'username': username, 'full_name': full_name}

            if not user_id:
                results['failed'] += 1
                user_info['reason'] = 'Отсутствует Telegram ID'
                results['failed_users'].append(user_info)
                continue

            success = False
            error_reason = 'Неизвестная ошибка'

            try:
                html_message = message if message else ""
                async with aiohttp.ClientSession() as session:
                    if files and len(files) > 0:
                        for idx, file_path in enumerate(files):
                            if os.path.exists(file_path):
                                url = f"{self.api_url}/sendDocument"
                                form_data = aiohttp.FormData()
                                form_data.add_field('chat_id', str(user_id))

                                if file_path in file_ids_map:
                                    form_data.add_field('document', file_ids_map[file_path])
                                    if idx == 0 and html_message:
                                        form_data.add_field('caption', html_message)
                                        form_data.add_field('parse_mode', 'HTML')

                                    async with session.post(url, data=form_data) as resp:
                                        res_json = await resp.json()
                                        if res_json.get('ok'):
                                            success = True
                                        else:
                                            error_reason = res_json.get('description', 'Telegram error')
                                else:
                                    with open(file_path, 'rb') as f:
                                        form_data.add_field('document', f, filename=os.path.basename(file_path))
                                        if idx == 0 and html_message:
                                            form_data.add_field('caption', html_message)
                                            form_data.add_field('parse_mode', 'HTML')

                                        async with session.post(url, data=form_data) as resp:
                                            res_json = await resp.json()
                                            if res_json.get('ok'):
                                                success = True
                                                doc = res_json.get('result', {}).get('document', {})
                                                if 'file_id' in doc:
                                                    file_ids_map[file_path] = doc['file_id']
                                            else:
                                                error_reason = res_json.get('description', 'Telegram error')
                                await asyncio.sleep(0.05)
                        if files and len(files) > 0 and success:
                            success = True
                    else:
                        url = f"{self.api_url}/sendMessage"
                        payload = {'chat_id': user_id, 'text': html_message, 'parse_mode': 'HTML'}
                        async with session.post(url, json=payload) as resp:
                            res_json = await resp.json()
                            success = res_json.get('ok', False)
                            if not success:
                                error_reason = res_json.get('description', 'Telegram error')

            except Exception as e:
                logger.error(f"Error sending to {user_id}: {e}")
                error_reason = str(e)
                success = False

            if success:
                results['success'] += 1
                results['success_users'].append(user_info)
            else:
                results['failed'] += 1
                user_info['reason'] = error_reason
                results['failed_users'].append(user_info)

        results['file_count'] = file_count
        results['file_ids'] = list(file_ids_map.values())
        return results

bot = TelegramBot(os.getenv('TG_BOT_TOKEN', ''))

# ============================================
# АВТОРИЗАЦИЯ И ОСНОВНЫЕ СТРАНИЦЫ
# ============================================

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')

        manager = run_async(db.verify_manager(username, password))
        if manager:
            session['logged_in'] = True
            session['username'] = manager['username']
            session['manager_id'] = manager['id']
            session['full_name'] = manager['full_name']
            session['role'] = manager['role']
            session['groups'] = manager['group_names']
            session['login_time'] = datetime.now().isoformat()
            flash(f'Добро пожаловать, {manager["full_name"] or manager["username"]}!', 'success')
            return redirect(url_for('dashboard'))

        admin = run_async(db.verify_admin(username, password))
        if admin:
            session['logged_in'] = True
            session['username'] = admin['username']
            session['manager_id'] = admin['id']
            session['full_name'] = admin['full_name']
            session['role'] = admin['role']
            session['groups'] = ['admin', 'pr', 'event', 'travel']
            session['login_time'] = datetime.now().isoformat()
            flash(f'Добро пожаловать, {admin["full_name"] or admin["username"]}!', 'success')
            return redirect(url_for('dashboard'))
        else:
            flash('Неверное имя пользователя или пароль', 'danger')

    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    flash('Вы вышли из системы', 'info')
    return redirect(url_for('login'))

@app.route('/')
@login_required
def dashboard():
    try:
        groups = session.get('groups', [])
        stats = run_async(db.get_stats())
        broadcasts = run_async(db.get_recent_broadcasts(5))
        conferences = run_async(db.get_conferences_list())
        active_conferences = conferences[:5] if conferences else []
        users = run_async(db.get_all_users_with_details())
        recent_users = sorted(users, key=lambda x: x.get('last_active') or datetime.min, reverse=True)[:5]

        return render_template('dashboard.html',
                               stats=stats,
                               groups=groups,
                               broadcasts=broadcasts,
                               active_conferences=active_conferences,
                               recent_users=recent_users,
                               username=session.get('username'),
                               full_name=session.get('full_name'))
    except Exception as e:
        print(f"Dashboard error: {e}")
        flash('Ошибка загрузки данных', 'danger')
        return render_template('dashboard.html',
                               stats={'total_users': 0, 'active_today': 0, 'companies_stats': [],
                                      'banner_requests': 0, 'visa_requests': 0, 'active_conferences': 0,
                                      'total_broadcasts': 0},
                               groups=[],
                               broadcasts=[],
                               active_conferences=[],
                               recent_users=[],
                               username=session.get('username'),
                               full_name=session.get('full_name'))

@app.route('/broadcast', methods=['GET', 'POST'])
@login_required
def broadcast():
    companies = run_async(db.get_all_companies_from_config()) or []
    conferences = run_async(db.get_conferences_list()) or []
    users = run_async(db.get_all_users_with_details()) or []

    if request.method == 'POST':
        message = request.form.get('message')
        target_type = request.form.get('target_type', 'all')
        files = request.files.getlist('files')
        saved_files = []

        for file in files:
            if file and file.filename:
                if not allowed_file(file.filename):
                    flash(f'Файл {file.filename} имеет запрещенный формат!', 'danger')
                    continue
                filename = secure_filename(file.filename)
                filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
                file.save(filepath)
                saved_files.append(filepath)

        if not message and not saved_files:
            flash('Введите сообщение или прикрепите файл для рассылки', 'danger')
            return render_template('broadcast.html', companies=companies, conferences=conferences, users=users)

        if message and len(message) > 4000:
            flash('Сообщение слишком длинное (макс. 4000 символов)', 'danger')
            return render_template('broadcast.html', companies=companies, conferences=conferences, users=users)

        users_to_send = []
        if target_type == 'all':
            users_to_send = run_async(db.get_users_by_company_list())
        elif target_type == 'company':
            companies_selected = request.form.getlist('companies')
            if not companies_selected:
                flash('Выберите хотя бы одну компанию', 'danger')
                return render_template('broadcast.html', companies=companies, conferences=conferences, users=users)
            users_to_send = run_async(db.get_users_by_company_list(companies_selected))
        elif target_type == 'specific':
            user_ids = request.form.getlist('selected_users')
            if not user_ids:
                flash('Выберите хотя бы одного пользователя', 'danger')
                return render_template('broadcast.html', companies=companies, conferences=conferences, users=users)
            user_ids = [int(uid) for uid in user_ids]
            users_to_send = run_async(db.get_users_by_ids_list(user_ids))
        elif target_type == 'conference':
            conferences_selected = request.form.getlist('conferences')
            if not conferences_selected:
                flash('Выберите хотя бы одну конференцию', 'danger')
                return render_template('broadcast.html', companies=companies, conferences=conferences, users=users)
            users_to_send = run_async(db.get_users_by_conference_list(conferences_selected))

        if not users_to_send:
            flash('Нет пользователей для рассылки', 'warning')
            return render_template('broadcast.html', companies=companies, conferences=conferences, users=users)

        results = run_async(bot.broadcast_to_users(users_to_send, message, saved_files))
        telegram_file_ids = results.get('file_ids', [])
        manager_name = session.get('full_name') or session.get('username') or 'Менеджер'
        display_name = f"{manager_name} (Рассылка)"
        messages_to_insert = []

        for user in users_to_send:
            uid = user.get('user_id')
            if uid:
                if saved_files and telegram_file_ids:
                    for idx, tg_file_id in enumerate(telegram_file_ids):
                        text_to_save = message if idx == 0 and message else ""
                        messages_to_insert.append(
                            (uid, display_name, 'outgoing', text_to_save, None, tg_file_id)
                        )
                elif message:
                    messages_to_insert.append(
                        (uid, display_name, 'outgoing', message, None, None)
                    )

        if messages_to_insert:
            run_async(db.save_user_messages_bulk(messages_to_insert))

        for filepath in saved_files:
            try:
                os.remove(filepath)
            except Exception:
                pass

        company_info = ', '.join(request.form.getlist('companies')) if target_type == 'company' else target_type
        run_async(db.log_broadcast(
            username=session.get('username', 'admin'),
            broadcast_type=target_type,
            company=company_info,
            success=results['success'],
            failed=results['failed'],
            message_length=len(message) if message else 0,
            file_count=results.get('file_count', 0),
            success_users=results.get('success_users', []),
            failed_users=results.get('failed_users', [])
        ))

        flash(f'Рассылка завершена! Успешно: {results["success"]}, Ошибок: {results["failed"]}, Файлов: {results.get("file_count", 0)}', 'success')
        return redirect(url_for('dashboard'))

    return render_template('broadcast.html', companies=companies, conferences=conferences, users=users)

@app.route('/conferences')
@login_required
def conferences_page():
    conferences = run_async(db.get_conferences_list())
    stats = {
        'total_conferences': len(conferences),
        'active_conferences': len([c for c in conferences if c.get('user_count', 0) > 0]),
        'upcoming_conferences': 0,
        'total_participants': sum(c.get('user_count', 0) for c in conferences)
    }
    return render_template('conferences.html', conferences=conferences, stats=stats, username=session.get('username'))

@app.route('/users')
@login_required
def users_page():
    try:
        users = run_async(db.get_all_users_with_details())
        companies = run_async(db.get_all_companies_from_config())
        conferences = run_async(db.get_conferences_list())
        stats = {
            'total': len(users),
            'online': len([u for u in users if u.get('is_online')]),
            'active_today': len([u for u in users if u.get('last_active') and u['last_active'] > datetime.now() - timedelta(days=1)]),
            'with_requests': len([u for u in users if u.get('requests_count', 0) > 0]),
            'companies': len(companies),
            'conferences': len(conferences)
        }
        return render_template('users.html', users=users, companies=companies, conferences=conferences, stats=stats, username=session.get('username'))
    except Exception as e:
        print(f"Users page error: {e}")
        flash('Ошибка загрузки пользователей', 'danger')
        return render_template('users.html', users=[], companies=[], conferences=[], stats={'total': 0, 'online': 0, 'active_today': 0, 'with_requests': 0, 'companies': 0, 'conferences': 0}, username=session.get('username'))

# ============================================
# API ДЛЯ УПРАВЛЕНИЯ ПОЛЬЗОВАТЕЛЯМИ
# ============================================

@app.route('/api/users/<int:user_id>', methods=['GET'])
@login_required
def api_get_user(user_id):
    user = run_async(db.get_user_details_by_id(user_id))
    if user:
        return jsonify(user)
    return jsonify({'error': 'User not found'}), 404

@app.route('/api/broadcast/<int:broadcast_id>')
@login_required
def api_get_broadcast_details(broadcast_id):
    details = run_async(db.get_broadcast_details_by_id(broadcast_id))
    if details:
        return jsonify(details)
    return jsonify({'error': 'Broadcast not found'}), 404

@app.route('/api/broadcasts/all')
@login_required
def get_all_broadcasts_api():
    broadcasts = run_async(db.get_recent_broadcasts(limit=500)) or []
    result = []
    for b in broadcasts:
        b_dict = dict(b)
        if b_dict.get('timestamp'):
            b_dict['formatted_time'] = b_dict['timestamp'].strftime('%d.%m.%Y %H:%M')
        result.append(b_dict)
    return jsonify(result)

@app.route('/api/users/<int:user_id>', methods=['PUT'])
@login_required
def api_update_user(user_id):
    data = request.get_json()
    if not data:
        return jsonify({'error': 'No data provided'}), 400

    user_data = {
        'user_id': user_id,
        'username': data.get('username'),
        'full_name': data.get('full_name'),
        'position': data.get('position'),
        'company': data.get('company'),
        'language': data.get('language', 'ru')
    }
    success = run_async(db.save_user_registration(user_data))
    if success:
        return jsonify({'success': True})
    return jsonify({'error': 'Failed to update user'}), 500

@app.route('/api/users/<int:user_id>', methods=['DELETE'])
@login_required
def api_delete_user(user_id):
    return jsonify({'success': True})

@app.route('/api/users/<int:user_id>/block', methods=['POST'])
@login_required
def api_block_user(user_id):
    return jsonify({'success': True})

@app.route('/api/users/<int:user_id>/unblock', methods=['POST'])
@login_required
def api_unblock_user(user_id):
    return jsonify({'success': True})

@app.route('/api/users/export', methods=['GET'])
@login_required
def api_export_users():
    users = run_async(db.get_all_users_with_details())
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['ID', 'Username', 'Full Name', 'Company', 'Position', 'Language', 'Registered', 'Last Active', 'Requests Count'])
    for user in users:
        writer.writerow([
            user['user_id'],
            user['username'],
            user.get('full_name', ''),
            user.get('company', ''),
            user.get('position', ''),
            user.get('language', 'ru'),
            user.get('registered_at', '').strftime('%Y-%m-%d %H:%M') if user.get('registered_at') else '',
            user.get('last_active', '').strftime('%Y-%m-%d %H:%M') if user.get('last_active') else '',
            user.get('requests_count', 0)
        ])
    output.seek(0)
    return send_file(
        io.BytesIO(output.getvalue().encode('utf-8-sig')),
        mimetype='text/csv',
        as_attachment=True,
        download_name=f'users_export_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv'
    )

@app.route('/api/users/count', methods=['GET', 'POST'])
@login_required
def api_users_count():
    if request.method == 'POST':
        data = request.get_json()
        companies = data.get('companies', [])
        conferences = data.get('conferences', [])
    else:
        companies = request.args.get('companies', '').split(',')
        conferences = request.args.get('conferences', '').split(',')
        if companies == ['']: companies = []
        if conferences == ['']: conferences = []

    if companies:
        users = run_async(db.get_users_by_company_list(companies))
    elif conferences:
        users = run_async(db.get_users_by_conference_list(conferences))
    else:
        users = run_async(db.get_users_by_company_list())
    return jsonify({'count': len(users)})

# ============================================
# ПАНЕЛИ И СПЕЦИФИЧЕСКИЕ API
# ============================================

@app.route('/event')
@login_required
@group_required(['event', 'admin'])
def event_panel():
    ticket_requests = run_async(db.get_all_ticket_requests())
    ticket_stats = run_async(db.get_ticket_request_stats())
    stands = run_async(db.get_all_event_stands()) or []
    companies = run_async(db.get_all_companies_from_config()) or []
    conferences = run_async(db.get_conferences_list()) or []

    return render_template(
        'event_panel.html',
        ticket_requests=ticket_requests,
        stats={'ticket_stats': ticket_stats},
        stands=stands,
        companies=companies,
        conferences=conferences,
        username=session.get('username')
    )
@app.route('/travel')
@login_required
@group_required(['travel', 'admin'])
def travel_panel():
    stats = run_async(db.get_travel_stats())
    requests = run_async(db.get_all_travel_flight_requests())
    return render_template('travel_panel.html', requests=requests, stats=stats, username=session.get('username'))

@app.route('/pr')
@login_required
@group_required(['pr', 'admin'])
def pr_panel():
    banner_requests = run_async(db.get_all_banner_requests())
    business_cards = run_async(db.get_all_business_cards())
    affiliate_bookings = run_async(db.get_all_affiliate_bookings()) or []
    affiliate_reports = run_async(db.get_all_affiliate_reports()) or []

    # Older bookings may contain Russian or English labels saved by the bot.
    # Normalize only known option values for this page; leave free text untouched.
    option_keys = {
        'карта': 'affiliate_value_card', 'card': 'affiliate_value_card',
        'наличные': 'affiliate_value_cash', 'cash': 'affiliate_value_cash',
        'обычный': 'affiliate_value_regular', 'regular': 'affiliate_value_regular',
        'не указано': 'affiliate_value_unspecified', 'not specified': 'affiliate_value_unspecified',
    }
    lang = session.get('lang', 'ru')
    affiliate_bookings = [
        {
            **booking,
            **{
                field: get_text(option_keys[str(booking[field]).strip().casefold()], lang)
                if booking.get(field) is not None and str(booking[field]).strip().casefold() in option_keys
                else booking.get(field)
                for field in ('payment_method', 'partnertype')
            },
        }
        for booking in affiliate_bookings
    ]

    for req in banner_requests:
        if 'status' not in req or not req.get('status'):
            req['status'] = 'pending'
    for card in business_cards:
        if 'status' not in card or not card.get('status'):
            card['status'] = 'pending'

    banner_filter = request.args.get('banner_filter', 'all')
    cards_filter = request.args.get('cards_filter', 'all')
    filtered_banners = [r for r in banner_requests if r.get('status') == banner_filter] if banner_filter != 'all' else banner_requests
    filtered_cards = [c for c in business_cards if c.get('status') == cards_filter] if cards_filter != 'all' else business_cards

    banner_stats = {
        'total': len(banner_requests),
        'pending': len([r for r in banner_requests if r.get('status') == 'pending']),
        'in_progress': len([r for r in banner_requests if r.get('status') == 'in_progress']),
        'ready': len([r for r in banner_requests if r.get('status') == 'ready'])
    }
    cards_stats = {
        'total': len(business_cards),
        'pending': len([c for c in business_cards if c.get('status') == 'pending']),
        'in_progress': len([c for c in business_cards if c.get('status') == 'in_progress']),
        'ready': len([c for c in business_cards if c.get('status') == 'ready'])
    }

    return render_template('pr_panel.html',
                           banner_requests=filtered_banners,
                           business_cards=filtered_cards,
                           affiliate_bookings=affiliate_bookings,
                           affiliate_reports=affiliate_reports,
                           banner_stats=banner_stats,
                           cards_stats=cards_stats,
                           banner_filter=banner_filter,
                           cards_filter=cards_filter,
                           username=session.get('username'))

@app.route('/admin_users')
@login_required
@role_required(['admin'])
def admin_users_page():
    admins = run_async(db.get_all_admin_users())
    return render_template('admin_users.html', admins=admins, username=session.get('username'))


@app.route('/admin_users/add', methods=['POST'])
@login_required
@role_required(['admin'])
def add_admin_user():
    username = (request.form.get('username') or '').strip()
    password = request.form.get('password') or ''
    full_name = (request.form.get('full_name') or '').strip() or None
    role = request.form.get('role', 'user')

    if not username or len(password) < 6 or role not in {'admin', 'user'}:
        flash('Проверьте имя пользователя, роль и пароль (минимум 6 символов)', 'danger')
        return redirect(url_for('admin_users_page'))

    permissions = {
        'manage_users': 'manage_users' in request.form,
        'broadcast': 'broadcast' in request.form,
        'view_stats': 'view_stats' in request.form,
        'manage_conferences': 'manage_conferences' in request.form,
    }
    success = run_async(db.add_admin_user(username, password, full_name, role, permissions))
    flash('Пользователь добавлен' if success else 'Не удалось добавить пользователя',
          'success' if success else 'danger')
    return redirect(url_for('admin_users_page'))


@app.route('/admin_users/edit/<int:admin_id>', methods=['POST'])
@login_required
@role_required(['admin'])
def edit_admin_user(admin_id):
    role = request.form.get('role', 'user')
    if role not in {'admin', 'user'}:
        flash('Недопустимая роль', 'danger')
        return redirect(url_for('admin_users_page'))

    data = {
        'full_name': (request.form.get('full_name') or '').strip() or None,
        'role': role,
        'can_manage_users': 'manage_users' in request.form,
        'can_broadcast': 'broadcast' in request.form,
        'can_view_stats': 'view_stats' in request.form,
        'can_manage_conferences': 'manage_conferences' in request.form,
    }
    password = request.form.get('password') or ''
    if password:
        if len(password) < 6:
            flash('Пароль должен содержать минимум 6 символов', 'danger')
            return redirect(url_for('admin_users_page'))
        data['password'] = password

    success = run_async(db.update_admin_user(admin_id, data))
    flash('Данные обновлены' if success else 'Не удалось обновить данные',
          'success' if success else 'danger')
    return redirect(url_for('admin_users_page'))


@app.route('/admin_users/delete/<int:admin_id>', methods=['POST'])
@login_required
@role_required(['admin'])
def delete_admin_user(admin_id):
    admins = run_async(db.get_all_admin_users()) or []
    target = next((admin for admin in admins if admin['id'] == admin_id), None)
    if not target:
        flash('Пользователь не найден', 'danger')
    elif target['username'] == session.get('username'):
        flash('Нельзя удалить текущую учётную запись', 'danger')
    else:
        success = run_async(db.delete_admin_user(admin_id))
        flash('Пользователь удалён' if success else 'Не удалось удалить пользователя',
              'success' if success else 'danger')
    return redirect(url_for('admin_users_page'))

@app.route('/change_password', methods=['GET', 'POST'])
@login_required
def change_password():
    if request.method == 'POST':
        old_password = request.form.get('old_password')
        new_password = request.form.get('new_password')
        confirm_password = request.form.get('confirm_password')

        if new_password != confirm_password:
            flash(get_text('passwords_not_match', session.get('lang', 'ru')), 'danger')
            return redirect(url_for('change_password'))

        if len(new_password) < 6:
            flash(get_text('password_min_length', session.get('lang', 'ru')), 'danger')
            return redirect(url_for('change_password'))

        manager = run_async(db.verify_manager(session['username'], old_password))
        if manager:
            success = run_async(db.update_manager_password(session['manager_id'], new_password))
            flash(get_text('password_changed_success', session.get('lang', 'ru')) if success else get_text('error_occurred', session.get('lang', 'ru')), 'success' if success else 'danger')
            return redirect(url_for('dashboard'))
        else:
            admin = run_async(db.verify_admin(session['username'], old_password))
            if admin:
                success = run_async(db.update_admin_user(admin['id'], {'password': new_password}))
                flash(get_text('password_changed_success', session.get('lang', 'ru')) if success else get_text('error_occurred', session.get('lang', 'ru')), 'success' if success else 'danger')
                return redirect(url_for('dashboard'))
            else:
                flash(get_text('invalid_current_password', session.get('lang', 'ru')), 'danger')

    return render_template('change_password.html')

@app.route('/api/per_diem/requests')
@login_required
def api_get_per_diem_requests():
    if session.get('role') != 'admin' and 'travel' not in session.get('groups', []):
        return jsonify({'error': 'Access denied'}), 403
    requests = run_async(db.get_all_per_diem_requests())
    return jsonify(requests)

@app.route('/api/per_diem/<int:request_id>/status', methods=['POST'])
@login_required
def api_update_per_diem_status(request_id):
    if session.get('role') != 'admin' and 'travel' not in session.get('groups', []):
        return jsonify({'error': 'Access denied'}), 403
    status = request.form.get('status')
    if status not in ['pending', 'in_progress', 'completed']:
        return jsonify({'error': 'Invalid status'}), 400
    success = run_async(db.update_per_diem_status(request_id, status))
    return jsonify({'success': success})

@app.route('/api/travel/requests')
@login_required
def api_get_travel_requests():
    if session.get('role') != 'admin' and 'travel' not in session.get('groups', []):
        return jsonify({'error': 'Access denied'}), 403
    try:
        requests = run_async(db.get_all_travel_flight_requests())
        return jsonify(requests)
    except Exception as e:
        logger.error(f"Error getting travel requests: {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/travel/request/<int:request_id>')
@login_required
def api_get_travel_request(request_id):
    if session.get('role') != 'admin' and 'travel' not in session.get('groups', []):
        return jsonify({'error': 'Access denied'}), 403
    try:
        request_data = run_async(db.get_travel_flight_request_by_id(request_id))
        if request_data:
            return jsonify(request_data)
        return jsonify({'error': 'Request not found'}), 404
    except Exception as e:
        logger.error(f"Error getting travel request {request_id}: {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/travel/request/<int:request_id>/visa_status', methods=['POST'])
@login_required
def api_update_travel_visa_status(request_id):
    if session.get('role') != 'admin' and 'travel' not in session.get('groups', []):
        return jsonify({'error': 'Access denied'}), 403
    status = request.form.get('status')
    if status not in ['pending', 'in_progress', 'ready']:
        return jsonify({'error': 'Invalid status'}), 400

    current_request = run_async(db.get_travel_flight_request_by_id(request_id))
    old_status = current_request.get('visa_request_status') if current_request else None
    success = run_async(db.update_travel_visa_request_status(request_id, status))

    if success and current_request and old_status != status:
        user_data = run_async(db.get_user_data(current_request.get('user_id')))
        lang = user_data.get('language', 'ru') if user_data else 'ru'
        run_async(notify_travel_visa_status_change(
            user_id=current_request.get('user_id'),
            request_id=request_id,
            old_status=old_status or 'pending',
            new_status=status,
            username=current_request.get('username'),
            lang=lang
        ))
        return jsonify({'success': True})
    return jsonify({'success': False}), 500

@app.route('/api/travel/request/<int:request_id>/flight_status', methods=['POST'])
@login_required
def api_update_travel_flight_status(request_id):
    if session.get('role') != 'admin' and 'travel' not in session.get('groups', []):
        return jsonify({'error': 'Access denied'}), 403
    status = request.form.get('status')
    if status not in ['pending', 'in_progress', 'purchased']:
        return jsonify({'error': 'Invalid status'}), 400

    current_request = run_async(db.get_travel_flight_request_by_id(request_id))
    old_status = current_request.get('flight_request_status') if current_request else None
    success = run_async(db.update_travel_flight_request_status(request_id, status))

    if success and current_request and old_status != status:
        user_data = run_async(db.get_user_data(current_request.get('user_id')))
        lang = user_data.get('language', 'ru') if user_data else 'ru'
        run_async(notify_travel_flight_status_change(
            user_id=current_request.get('user_id'),
            request_id=request_id,
            old_status=old_status or 'pending',
            new_status=status,
            username=current_request.get('username'),
            lang=lang
        ))
        return jsonify({'success': True})
    return jsonify({'success': False}), 500

@app.route('/set_language/<lang>')
def set_language(lang):
    if lang in ['ru', 'en']:
        session['lang'] = lang
    return redirect(request.referrer or url_for('dashboard'))

@app.context_processor
def utility_processor():
    def _get_text(key, default=None, **kwargs):
        lang = session.get('lang', 'ru')
        text = get_text(key, lang)
        if text == key and default is not None:
            text = default
        if kwargs:
            try:
                text = text.format(**kwargs)
            except (KeyError, IndexError):
                pass
        return text
    return dict(get_text=_get_text)

@app.route('/api/visa/<int:request_id>/status', methods=['POST'])
@login_required
def update_visa_status(request_id):
    status = request.form.get('status')
    success = run_async(db.update_visa_status(request_id, status))
    if success:
        flash(f'Статус заявки #{request_id} обновлен на "{status}"', 'success')
    else:
        flash('Ошибка при обновлении статуса', 'danger')
    return redirect(url_for('travel_panel'))

@app.route('/api/banner/<int:request_id>/status', methods=['POST'])
@login_required
def update_banner_status(request_id):
    status = request.form.get('status')
    if status not in ['pending', 'in_progress', 'ready']:
        return jsonify({'success': False, 'error': 'Invalid status'}), 400

    current_request = run_async(db.get_banner_request_by_id(request_id))
    old_status = current_request.get('status') if current_request else None
    success = run_async(db.update_banner_status(request_id, status))

    if success and current_request and old_status != status:
        user_data = run_async(db.get_user_data(current_request.get('user_id')))
        lang = user_data.get('language', 'ru') if user_data else 'ru'
        run_async(notify_banner_status_change(
            user_id=current_request.get('user_id'),
            request_id=request_id,
            old_status=old_status or 'pending',
            new_status=status,
            username=current_request.get('username'),
            lang=lang
        ))
        return jsonify({'success': True})
    return jsonify({'success': False}), 500

@app.route('/api/business_card/<int:request_id>/status', methods=['POST'])
@login_required
def update_business_card_status(request_id):
    status = request.form.get('status')
    if status not in ['pending', 'in_progress', 'ready']:
        return jsonify({'success': False, 'error': 'Invalid status'}), 400

    current_request = run_async(db.get_business_card_request_by_id(request_id))
    old_status = current_request.get('status') if current_request else None
    success = run_async(db.update_business_card_status(request_id, status))

    if success and current_request and old_status != status:
        user_data = run_async(db.get_user_data(current_request.get('user_id')))
        lang = user_data.get('language', 'ru') if user_data else 'ru'
        run_async(notify_business_card_status_change(
            user_id=current_request.get('user_id'),
            request_id=request_id,
            old_status=old_status or 'pending',
            new_status=status,
            username=current_request.get('username'),
            lang=lang
        ))
        return jsonify({'success': True})
    return jsonify({'success': False}), 500

@app.route('/api/users/all')
@login_required
def api_get_all_users():
    users = run_async(db.get_all_users_basic())
    return jsonify(users)

@app.route('/api/users/by_conferences', methods=['POST'])
@login_required
def api_get_users_by_conferences():
    data = request.get_json()
    conferences = data.get('conferences', [])
    if not conferences:
        return jsonify([])
    users = run_async(db.get_users_by_conference_list(conferences))
    return jsonify(users)

@app.route('/api/visa/<int:request_id>')
@login_required
def api_visa_details(request_id):
    data = run_async(db.get_travel_flight_request_by_id(request_id))
    return jsonify(data)

@app.route('/api/banner/<int:request_id>')
@login_required
def api_banner_details(request_id):
    data = run_async(db.get_banner_request_by_id(request_id))
    return jsonify(data)

@app.route('/api/business_card/<int:request_id>')
@login_required
def api_business_card_details(request_id):
    data = run_async(db.get_business_card_request_by_id(request_id))
    return jsonify(data)

@app.route('/api/conferences/<path:conference_name>')
@login_required
def get_conference_details(conference_name):
    if not session.get('logged_in'):
        return jsonify({'error': 'Unauthorized'}), 401
    decoded_name = unquote(conference_name)

    try:
        def _get_details():
            async def fetch():
                async with db.pool.acquire() as conn:
                    conference = await conn.fetchrow(f"""
                        SELECT conference_name as name, city, start_date, end_date, bot_link, additional_info, is_active
                        FROM {db.db_schema_config}.conferences
                        WHERE conference_name = $1
                    """, decoded_name)

                    if not conference:
                        conference = await conn.fetchrow(f"""
                            SELECT conference_name as name, city, start_date, end_date, bot_link, additional_info, is_active
                            FROM {db.db_schema_config}.conferences
                            WHERE LOWER(conference_name) = LOWER($1)
                        """, decoded_name)

                    if not conference:
                        return None

                    users = await conn.fetch(f"""
                        SELECT DISTINCT up.user_id, up.username, up.full_name, up.company, up.position
                        FROM {db.db_schema_config}.user_profiles up
                        JOIN {db.db_schema}.user_conferences uc ON up.username = uc.username
                        WHERE uc.conference_name = $1
                        ORDER BY up.username
                    """, decoded_name)

                    result = dict(conference)
                    result['user_count'] = len(users) if users else 0
                    result['users'] = [dict(u) for u in users] if users else []
                    return result

            return run_async(fetch())

        result = _get_details()
        if result is None:
            return jsonify({'error': f'Conference "{decoded_name}" not found'}), 404
        return jsonify(result)
    except Exception as e:
        logger.error(f"Error getting conference details for '{decoded_name}': {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/stands/save', methods=['POST'])
@login_required
@group_required(['event', 'admin'])
def api_save_stand():
    company = request.form.get('company', '').strip()
    conference = request.form.get('conference', '').strip()

    if not company or not conference:
        flash('Компания и конференция обязательны для заполнения', 'danger')
        return redirect(url_for('event_panel'))

    photo = request.files.get('photo')
    saved_filename = None

    if photo and photo.filename:
        if allowed_image(photo.filename):
            ext = photo.filename.rsplit('.', 1)[1].lower()
            saved_filename = f"stand_{uuid.uuid4().hex}.{ext}"
            photo_save_path = os.path.join(STANDS_UPLOAD_FOLDER, saved_filename)
            photo.save(photo_save_path)
        else:
            flash('Недопустимый формат изображения (разрешены png, jpg, jpeg, webp)', 'danger')
            return redirect(url_for('event_panel'))

    stand_data = {
        'company': company,
        'conference': conference,
        'stand_style': request.form.get('stand_style', '').strip(),
        'stand_number': request.form.get('stand_number', '').strip(),
        'working_hours': request.form.get('working_hours', '').strip(),
        'dress_code': request.form.get('dress_code', '').strip(),
        'photo_path': saved_filename
    }

    success = run_async(db.upsert_event_stand(stand_data))
    if success:
        flash('Информация о стенде сохранена', 'success')
    else:
        flash('Ошибка при сохранении стенда', 'danger')

    return redirect(url_for('event_panel'))

@app.route('/api/stands/delete/<int:stand_id>', methods=['POST'])
@login_required
@group_required(['event', 'admin'])
def api_delete_stand(stand_id):
    photo_file = run_async(db.delete_event_stand(stand_id))
    if photo_file:
        full_path = os.path.join(STANDS_UPLOAD_FOLDER, photo_file)
        if os.path.exists(full_path):
            try:
                os.remove(full_path)
            except Exception as e:
                logger.error(f"Error removing stand photo file {photo_file}: {e}")
    flash('Стенд успешно удален', 'success')
    return redirect(url_for('event_panel'))

@app.route('/admin_managers')
@login_required
@group_required(['admin'])
def admin_managers():
    if session.get('role') != 'admin':
        flash('Доступ запрещен', 'danger')
        return redirect(url_for('dashboard'))
    managers = run_async(db.get_all_managers())
    groups = run_async(db.get_manager_groups())
    return render_template('admin_managers.html', managers=managers, groups=groups)

@app.route('/admin_managers/add', methods=['POST'])
@login_required
def add_manager():
    if session.get('role') != 'admin':
        return redirect(url_for('dashboard'))
    username = request.form.get('username')
    password = request.form.get('password')
    full_name = request.form.get('full_name')
    groups = request.form.getlist('groups')
    success = run_async(db.add_manager(username, password, full_name, groups))
    flash('Менеджер добавлен' if success else 'Ошибка при добавлении', 'success' if success else 'danger')
    return redirect(url_for('admin_managers'))

@app.route('/admin_managers/delete/<int:manager_id>', methods=['POST'])
@login_required
def delete_manager(manager_id):
    if session.get('role') != 'admin':
        flash('Доступ запрещен', 'danger')
        return redirect(url_for('admin_managers'))
    success = run_async(db.delete_manager(manager_id))
    flash('Менеджер успешно удален' if success else 'Ошибка при удалении. Невозможно удалить главного администратора.', 'success' if success else 'danger')
    return redirect(url_for('admin_managers'))

@app.route('/admin_managers/<int:manager_id>')
@login_required
@group_required(['admin'])
def get_manager_json(manager_id):
    managers = run_async(db.get_all_managers())
    manager = next((m for m in managers if m['id'] == manager_id), None)
    if manager:
        return jsonify(manager)
    return jsonify({'error': 'Not found'}), 404

@app.route('/admin_managers/edit/<int:manager_id>', methods=['POST'])
@login_required
@group_required(['admin'])
def edit_manager(manager_id):
    full_name = request.form.get('full_name')
    is_active = request.form.get('is_active') == 'true'
    groups = request.form.getlist('groups')
    success = run_async(db.update_manager(manager_id, full_name, is_active, groups))
    flash('Данные менеджера обновлены' if success else 'Ошибка при обновлении', 'success' if success else 'danger')
    return redirect(url_for('admin_managers'))

@app.route('/admin_managers/reset_password/<int:manager_id>', methods=['POST'])
@login_required
@group_required(['admin'])
def reset_manager_password(manager_id):
    new_password = request.form.get('new_password')
    confirm_password = request.form.get('confirm_password')
    if new_password != confirm_password:
        flash('Пароли не совпадают', 'danger')
        return redirect(url_for('admin_managers'))
    if len(new_password) < 6:
        flash('Пароль должен содержать минимум 6 символов', 'danger')
        return redirect(url_for('admin_managers'))
    success = run_async(db.update_manager_password(manager_id, new_password))
    flash('Пароль успешно изменен' if success else 'Ошибка при смене пароля', 'success' if success else 'danger')
    return redirect(url_for('admin_managers'))

# ============================================
# ЧАТЫ И ПАПКИ МЕНЕДЖЕРА
# ============================================

@app.route('/user_chats')
@login_required
def user_chats():
    manager_id = session.get('manager_id', 0)
    conversations = run_async(db.get_user_conversations(manager_id=manager_id))
    folders = run_async(db.get_manager_folders(manager_id=manager_id)) or []
    return render_template('user_chats.html', conversations=conversations, folders=folders)

@app.route('/api/conversations')
@login_required
def api_get_conversations():
    manager_id = session.get('manager_id', 0)
    conversations = run_async(db.get_user_conversations(manager_id=manager_id))
    for conv in conversations:
        if conv.get('last_message_time'):
            conv['last_message_time'] = conv['last_message_time'].isoformat()
    return jsonify(conversations)

@app.route('/api/chat/<int:user_id>/toggle_unread', methods=['POST'])
@login_required
def api_toggle_chat_unread(user_id):
    manager_id = session.get('manager_id', 0)
    is_unread = run_async(db.toggle_chat_unread(manager_id, user_id))
    return jsonify({'success': True, 'is_unread': is_unread})

@app.route('/api/chat/<int:user_id>/toggle_archive', methods=['POST'])
@login_required
def api_toggle_chat_archive(user_id):
    manager_id = session.get('manager_id', 0)
    is_archived = run_async(db.toggle_chat_archive(manager_id, user_id))
    return jsonify({'success': True, 'is_archived': is_archived})

@app.route('/api/chat/<int:user_id>/toggle_mute', methods=['POST'])
@login_required
def api_toggle_chat_mute(user_id):
    manager_id = session.get('manager_id', 0)
    is_muted = run_async(db.toggle_chat_mute(manager_id, user_id))
    return jsonify({'success': True, 'is_muted': is_muted})


@app.route('/api/folders', methods=['GET'])
@login_required
def api_get_folders():
    manager_id = session.get('manager_id', 0)
    folders = run_async(db.get_manager_folders(manager_id=manager_id))
    return jsonify(folders or [])

@app.route('/api/folders', methods=['POST'])
@login_required
def api_create_folder():
    manager_id = session.get('manager_id', 0)
    data = request.get_json() or {}
    name = (data.get('name') or '').strip()
    if not name:
        return jsonify({'error': 'Название папки не может быть пустым'}), 400
    folder = run_async(db.create_manager_folder(manager_id, name))
    if folder:
        return jsonify(folder)
    return jsonify({'error': 'Не удалось создать папку'}), 500

@app.route('/api/folders/<int:folder_id>', methods=['DELETE'])
@login_required
def api_delete_folder(folder_id):
    manager_id = session.get('manager_id', 0)
    success = run_async(db.delete_manager_folder(manager_id, folder_id))
    return jsonify({'success': success})

@app.route('/api/folders/<int:folder_id>/toggle', methods=['POST'])
@login_required
def api_toggle_folder_user(folder_id):
    manager_id = session.get('manager_id', 0)
    data = request.get_json() or {}
    user_id = data.get('user_id')
    if not user_id:
        return jsonify({'error': 'user_id required'}), 400
    res = run_async(db.toggle_user_in_folder(manager_id, folder_id, int(user_id)))
    return jsonify(res)

@app.route('/api/user_folders/<int:user_id>', methods=['GET'])
@login_required
def api_get_user_folders(user_id):
    manager_id = session.get('manager_id', 0)
    folders_map = run_async(db.get_manager_user_folders_map(manager_id)) or {}
    return jsonify({'user_id': user_id, 'folder_ids': folders_map.get(user_id, [])})

@app.route('/api/questions/forward', methods=['POST'])
@login_required
def api_forward_question():
    data = request.get_json()
    question_id = data.get('question_id')
    source_department = data.get('source_department')
    target_department = data.get('target_department')
    source_table = data.get('source_table')

    if not all([question_id, source_department, target_department, source_table]):
        return jsonify({'error': 'Missing required fields'}), 400

    manager_groups = session.get('groups', [])
    role = session.get('role')
    if role != 'admin' and source_department not in manager_groups:
        return jsonify({'error': f'Access denied to {source_department} department'}), 403

    question = run_async(db.get_question_by_id(source_table, question_id))
    if not question:
        return jsonify({'error': 'Question not found'}), 404

    target_tables = {
        'pr': f'{db.db_schema_pr}.pr_questions',
        'event': f'{db.db_schema_event}.event_questions',
        'travel': f'{db.db_schema_travel}.travel_questions'
    }
    target_table = target_tables.get(target_department)
    if not target_table:
        return jsonify({'error': 'Invalid target department'}), 400

    forwarded_question = {
        'username': question.get('username'),
        'user_id': question.get('user_id'),
        'category': f"forwarded_from_{source_department}",
        'question': f"[🔄 Переслано из отдела {source_department.upper()} менеджером @{session.get('username')}]\n\n📝 Оригинальный вопрос:\n{question.get('question')}",
        'created_at': question.get('created_at')
    }

    success = run_async(db.save_forwarded_question(target_table, forwarded_question))
    if success:
        run_async(db.log_user_action(
            user_id=session.get('manager_id', 0),
            username=session.get('username'),
            action="question_forwarded",
            details={
                "question_id": question_id,
                "source": source_department,
                "target": target_department,
                "source_table": source_table
            }
        ))
        return jsonify({'success': True, 'message': f'Question forwarded to {target_department.upper()}'})
    return jsonify({'error': 'Failed to forward question'}), 500

@app.route('/api/questions/<int:question_id>/departments', methods=['GET'])
@login_required
def api_get_available_departments_for_forward(question_id):
    manager_groups = session.get('groups', [])
    role = session.get('role')
    all_departments = [
        {'id': 'pr', 'name': '📢 PR отдел'},
        {'id': 'event', 'name': '🎪 Event отдел'},
        {'id': 'travel', 'name': '✈️ Travel отдел'}
    ]
    source_department = request.args.get('source', '')
    if role == 'admin':
        available = [d for d in all_departments if d['id'] != source_department]
    else:
        available = [d for d in all_departments if d['id'] != source_department and d['id'] in manager_groups]
    return jsonify({'departments': available})

@app.route('/api/user_messages/<int:user_id>')
@login_required
def api_get_user_messages(user_id):
    messages = run_async(db.get_user_messages(user_id, limit=100))
    user_data = run_async(db.get_user_data(user_id))

    formatted_messages = []
    for msg in messages:
        formatted_messages.append({
            'id': msg['id'],
            'uid': f"{msg.get('source_type', 'msg')}_{msg['id']}",
            'source_type': msg.get('source_type', 'message'),
            'direction': msg['direction'],
            'message_text': msg['message_text'],
            'sender_username': msg.get('username', ''),
            'file_type': msg.get('file_type'),
            'file_id': msg.get('file_id'),
            'created_at': msg['created_at'].isoformat()
        })

    return jsonify({
        'user_id': user_id,
        'username': user_data.get('username') if user_data else str(user_id),
        'full_name': user_data.get('full_name', ''),
        'messages': formatted_messages
    })

@app.route('/api/user_messages/<int:user_id>/read', methods=['POST'])
@login_required
def api_mark_messages_read(user_id):
    success = run_async(db.mark_messages_read(user_id, session.get('manager_id')))
    return jsonify({'success': success})

@app.route('/admin/send_message_to_user', methods=['POST'])
@login_required
def send_message_to_user():
    user_id = request.form.get('user_id')
    message_text = request.form.get('message_text')
    uploaded_file = request.files.get('file')

    if not user_id:
        flash(get_text('error_user_not_found'), "danger")
        return redirect(url_for('travel_panel'))

    manager_username = session.get('username', 'Manager')
    manager_full_name = session.get('full_name', 'Manager')
    final_text = f"{message_text}"

    filepath = None
    if uploaded_file and uploaded_file.filename:
        filename = secure_filename(uploaded_file.filename)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        uploaded_file.save(filepath)

    try:
        tg_res = run_async(bot.send_message(chat_id=user_id, text=final_text, file=filepath))
        if tg_res:
            file_id_to_save = tg_res if isinstance(tg_res, str) else None
            manager_display_name = f"{manager_full_name} (@{manager_username})"
            run_async(db.save_user_message(
                user_id=int(user_id),
                username=manager_display_name,
                message_text=final_text,
                file_type=uploaded_file.content_type if uploaded_file and uploaded_file.filename else None,
                file_id=file_id_to_save,
                direction='outgoing'
            ))
            flash(get_text('message_sent_success'), "success")
        else:
            flash(get_text('error_sending_tg'), "danger")
    except Exception as e:
        logger.error(f"Error sending message: {e}")
        flash(f"{get_text('error_occurred')}: {e}", "danger")
    finally:
        if filepath and os.path.exists(filepath):
            os.remove(filepath)

    return redirect(url_for('travel_panel'))

@app.route('/api/send_message', methods=['POST'])
@login_required
def api_send_message():
    user_id = request.form.get('user_id')
    message = request.form.get('message', '')
    if not user_id:
        return jsonify({'success': False, 'error': 'User ID required'}), 400

    files = request.files.getlist('files')
    if not files and 'file' in request.files:
        files = [request.files['file']]

    if len(files) > 10:
        return jsonify({'success': False, 'error': 'Максимум 10 файлов за раз'}), 400

    saved_files = []
    for file in files:
        if file and file.filename:
            if not allowed_file(file.filename):
                return jsonify({'success': False, 'error': f'Запрещенный формат: {file.filename}'}), 400
            filename = secure_filename(file.filename)
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            file.save(filepath)
            saved_files.append(filepath)

    try:
        sender_name = session.get('full_name') or session.get('username') or 'Менеджер'

        if saved_files:
            tg_file_ids = run_async(bot.send_documents(user_id, message, saved_files), timeout=300)
            if not tg_file_ids or len(tg_file_ids) != len(saved_files):
                logger.error("Telegram sent %s/%s files to user %s", len(tg_file_ids or []), len(saved_files), user_id)
                return jsonify({'success': False, 'error': 'Telegram не подтвердил отправку всех файлов. Проверьте чат перед повторной отправкой.'}), 502
            messages_to_insert = []
            if tg_file_ids:
                for idx, tg_file_id in enumerate(tg_file_ids):
                    text_to_save = message if idx == 0 else ""
                    messages_to_insert.append(
                        (int(user_id), sender_name, 'outgoing', text_to_save, None, tg_file_id)
                    )
            if messages_to_insert:
                run_async(db.save_user_messages_bulk(messages_to_insert), timeout=60)
        else:
            res = run_async(bot.send_message(user_id, message), timeout=60)
            if res:
                run_async(db.save_user_messages_bulk([
                    (int(user_id), sender_name, 'outgoing', message, None, None)
                ]), timeout=60)
            else:
                return jsonify({'success': False, 'error': 'Не удалось отправить сообщение'}), 500

        return jsonify({'success': True})
    except Exception as e:
        logger.error(f"Error sending message in chat: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        for fp in saved_files:
            try:
                if os.path.exists(fp):
                    os.remove(fp)
            except Exception:
                pass

@app.route('/api/file/<file_id>')
@login_required
def api_get_file(file_id):
    bot_token = os.getenv('TG_BOT_TOKEN')
    if not bot_token or not file_id:
        return "Ошибка конфигурации", 500
    if len(file_id) < 20:
        return "Этот файл был удален с сервера и не имеет ID в Telegram", 404

    try:
        file_info_response = requests.get(
            f"https://api.telegram.org/bot{bot_token}/getFile?file_id={file_id}",
            timeout=20,
        )
        file_info_response.raise_for_status()
        file_info = file_info_response.json()
        if not file_info.get('ok'):
            return "Файл не найден в Telegram", 404

        file_path = file_info['result']['file_path']
        download_url = f"https://api.telegram.org/file/bot{bot_token}/{file_path}"
        filename = file_path.split('/')[-1]

        mime_type, _ = mimetypes.guess_type(filename)
        if not mime_type:
            ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
            if ext in {'jpg', 'jpeg', 'png', 'webp', 'gif'}:
                mime_type = f'image/{ext if ext != "jpg" else "jpeg"}'
            else:
                mime_type = 'application/octet-stream'

        file_data = requests.get(download_url, timeout=60)
        file_data.raise_for_status()
        force_download = request.args.get('download') == '1'

        response = send_file(
            io.BytesIO(file_data.content),
            mimetype=mime_type,
            as_attachment=force_download,
            download_name=filename
        )
        response.headers['Cache-Control'] = 'public, max-age=86400'
        return response
    except Exception as e:
        logger.error(f"Download/View error: {e}")
        return "Ошибка при загрузке файла", 500


@app.route('/api/tickets/export', methods=['GET'])
@login_required
@group_required(['event', 'admin'])
def api_export_tickets():
    ticket_requests = run_async(db.get_all_ticket_requests())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заявки на билеты"

    # Включаем отображение сетки таблицы
    ws.views.sheetView[0].showGridLines = True

    # Заголовки таблицы
    headers = [
        'ID', 'Telegram Username', 'ФИО', 'Должность',
        'Компания', 'Email', 'Телефон', 'Страна',
        'Дата создания', 'Статус'
    ]
    ws.append(headers)

    # Оформление шапки таблицы
    header_fill = PatternFill(start_color="3B5998", end_color="3B5998", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.row_dimensions[1].height = 26
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Соответствие статусов и их оформление
    status_titles = {
        'pending': 'Ожидание',
        'in_progress': 'В процессе',
        'ready': 'Готово'
    }
    status_fills = {
        'pending': PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
        'in_progress': PatternFill(start_color="CFE2FF", end_color="CFE2FF", fill_type="solid"),
        'ready': PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
    }
    status_fonts = {
        'pending': Font(name="Calibri", size=10, bold=True, color="856404"),
        'in_progress': Font(name="Calibri", size=10, bold=True, color="084298"),
        'ready': Font(name="Calibri", size=10, bold=True, color="155724"),
    }

    data_font = Font(name="Calibri", size=10)

    # Заполнение строк данными
    for row_idx, req in enumerate(ticket_requests, start=2):
        status_key = req.get('status')
        status_label = status_titles.get(status_key, status_key or '-')
        created_at_str = req['created_at'].strftime('%d.%m.%Y %H:%M') if req.get('created_at') else '-'

        row_values = [
            req.get('id', ''),
            f"@{req.get('username', '')}" if req.get('username') else '',
            req.get('full_name') or '-',
            req.get('position') or '-',
            req.get('company') or '-',
            req.get('email') or '-',
            req.get('phone') or '-',
            req.get('country') or '-',
            created_at_str,
            status_label
        ]
        ws.append(row_values)
        ws.row_dimensions[row_idx].height = 20

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border

            # Центрируем ID, дату и статус
            if col_idx in (1, 9, 10):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Цветной бейдж для статуса
            if col_idx == 10 and status_key in status_fills:
                cell.fill = status_fills[status_key]
                cell.font = status_fonts[status_key]

    # Автоматическая настройка ширины столбцов под текст
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f'ticket_requests_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

@app.route('/api/questions')
@login_required
def api_get_my_questions():
    manager_groups = session.get('groups', [])
    role = session.get('role')
    if role == 'admin':
        manager_groups = ['pr', 'event', 'travel']

    all_questions = []
    tables_map = {
        'pr': f'{db.db_schema_pr}.pr_questions',
        'event': f'{db.db_schema_event}.event_questions',
        'travel': f'{db.db_schema_travel}.travel_questions'
    }

    for group in manager_groups:
        if group not in tables_map:
            continue
        questions = run_async(db.get_all_questions_by_table(tables_map[group]))
        for q in questions:
            q['department'] = group
            all_questions.append(q)

    all_questions.sort(key=lambda x: x.get('created_at', datetime.min), reverse=True)
    return jsonify(all_questions)

@app.route('/api/affiliate/bookings')
@login_required
def api_affiliate_bookings():
    bookings = run_async(db.get_all_affiliate_bookings()) or []
    formatted = []
    for b in bookings:
        item = dict(b)
        if item.get('created_at'):
            item['created_at'] = item['created_at'].strftime('%d.%m.%Y %H:%M')
        if item.get('updated_at'):
            item['updated_at'] = item['updated_at'].strftime('%d.%m.%Y %H:%M')
        formatted.append(item)
    return jsonify(formatted)

@app.route('/api/affiliate/reports')
@login_required
def api_affiliate_reports():
    reports = run_async(db.get_all_affiliate_reports()) or []
    formatted = []
    for r in reports:
        item = dict(r)
        if item.get('created_at'):
            item['created_at'] = item['created_at'].strftime('%d.%m.%Y %H:%M')
        if item.get('updated_at'):
            item['updated_at'] = item['updated_at'].strftime('%d.%m.%Y %H:%M')
        formatted.append(item)
    return jsonify(formatted)

@app.route('/api/affiliate/booking/<int:booking_id>/status', methods=['POST'])
@login_required
def api_update_booking_status(booking_id):
    status = request.form.get('status')
    success = run_async(db.update_affiliate_booking_status(booking_id, status))
    return jsonify({'success': success})

@app.route('/api/affiliate/report/<int:report_id>/status', methods=['POST'])
@login_required
def api_update_report_status(report_id):
    status = request.form.get('status')
    success = run_async(db.update_affiliate_report_status(report_id, status))
    return jsonify({'success': success})

@app.route('/api/ticket/<int:request_id>/status', methods=['POST'])
@login_required
def api_update_ticket_status(request_id):
    user_groups = session.get('groups', [])
    role = session.get('role')
    if role != 'admin' and 'event' not in user_groups:
        return jsonify({'error': 'Access denied'}), 403

    status = request.form.get('status')
    if status not in ['pending', 'in_progress', 'ready']:
        return jsonify({'error': 'Invalid status'}), 400

    current_request = run_async(db.get_ticket_request_by_id(request_id))
    old_status = current_request.get('status') if current_request else None
    success = run_async(db.update_ticket_request_status(request_id, status))

    if success:
        if current_request:
            run_async(notify_ticket_status_change(
                user_id=current_request.get('user_id'),
                request_id=request_id,
                old_status=old_status or 'pending',
                new_status=status,
                username=current_request.get('username')
            ))
        run_async(db.log_user_action(
            user_id=session.get('manager_id', 0),
            username=session.get('username'),
            action="ticket_status_changed",
            details={"request_id": request_id, "new_status": status, "old_status": old_status}
        ))
        return jsonify({'success': True})
    return jsonify({'error': 'Failed to update status'}), 500

@app.route('/api/questions/share', methods=['POST'])
@login_required
def api_share_question():
    data = request.get_json()
    question_id = data.get('question_id')
    source_department = data.get('source_department')
    target_department = data.get('target_department')

    if not all([question_id, source_department, target_department]):
        return jsonify({'error': 'Missing required fields'}), 400

    manager_groups = session.get('groups', [])
    if source_department not in manager_groups and session.get('role') != 'admin':
        return jsonify({'error': 'Access denied'}), 403

    success = run_async(db.share_question_with_department(
        question_id=question_id,
        question_type=f"{source_department}_question",
        source_department=source_department,
        target_department=target_department,
        shared_by=session.get('username')
    ))
    return jsonify({'success': success})

# ============================================
# ЗАПУСК ПРИЛОЖЕНИЯ
# ============================================

if __name__ == '__main__':
    print("Initializing database...")
    try:
        init_db()
        print("✅ Database connection established")
    except Exception as e:
        print(f"❌ Database connection failed: {e}")
        print("Check your .env file and make sure PostgreSQL is running")
        print("⚠️ Continuing without database connection...")

    debug_enabled = os.getenv('FLASK_DEBUG', '').lower() in {'1', 'true', 'yes'}
    app.run(host='0.0.0.0', port=5005, debug=debug_enabled)
